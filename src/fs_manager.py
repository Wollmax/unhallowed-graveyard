"""Модуль файловой системы для UniVault.
Отвечает за работу с локальными директориями, валидацию и санитизацию путей,
создание структуры папок (01_Task, 02_Workspace, 03_Submission), открытие в проводнике/IDE.
Все пути обрабатываются через pathlib.Path.
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Dict, List, Optional


# Запрещенные символы в именах файлов и папок Windows: \ / : * ? " < > |
WINDOWS_INVALID_CHARS_REGEX = re.compile(r'[\\/:*?"<>|]')


def sanitize_folder_name(name: str, max_length: int = 60) -> str:
    """Очищает строку от символов, запрещенных в файловой системе Windows."""
    cleaned = WINDOWS_INVALID_CHARS_REGEX.sub("_", name.strip())
    # Убираем повторяющиеся подчеркивания и пробелы
    cleaned = re.sub(r"\s+", " ", cleaned)
    cleaned = re.sub(r"_+", "_", cleaned)
    # Убираем точки и пробелы на концах (Windows не разрешает)
    cleaned = cleaned.strip(". ")
    if not cleaned:
        cleaned = "untitled"
    return cleaned[:max_length]


class FSManager:
    """Менеджер файловой структуры рабочей области UniVault."""

    SUBFOLDER_TASK = "01_Task"
    SUBFOLDER_WORKSPACE = "02_Workspace"
    SUBFOLDER_SUBMISSION = "03_Submission"
    FOLDER_TEMPLATES = "_Templates"
    FOLDER_MATERIALS = "_Materials"

    def __init__(self, root_path: Path | str) -> None:
        self.root_path = Path(root_path).resolve()
        self.ensure_workspace()

    def ensure_workspace(self) -> Path:
        """Создает корневую папку и подпапку глобальных шаблонов _Templates."""
        self.root_path.mkdir(parents=True, exist_ok=True)
        templates_dir = self.root_path / self.FOLDER_TEMPLATES
        templates_dir.mkdir(parents=True, exist_ok=True)
        return self.root_path

    # --- Управление папками предметов ---

    def create_subject_folder(self, subject_name: str, short_name: Optional[str] = None) -> str:
        """
        Создает папку предмета и внутреннюю папку _Materials.
        Возвращает относительное имя папки (folder_path для базы данных).
        """
        tag = f"_{sanitize_folder_name(short_name, 15)}" if short_name else ""
        folder_name = sanitize_folder_name(f"{subject_name}{tag}")
        subject_dir = self.root_path / folder_name
        subject_dir.mkdir(parents=True, exist_ok=True)

        # Создаем папку общих материалов предмета
        materials_dir = subject_dir / self.FOLDER_MATERIALS
        materials_dir.mkdir(parents=True, exist_ok=True)

        return folder_name

    def get_subject_dir(self, subject_folder: str) -> Path:
        return self.root_path / subject_folder

    # --- Управление папками практических работ ---

    def create_task_folders(
        self,
        subject_folder: str,
        number: int,
        task_type: str,
        title: str,
    ) -> str:
        """
        Создает папку практической работы со структурой:
          [Task_Title]/
            ├── 01_Task/
            ├── 02_Workspace/
            └── 03_Submission/
        Возвращает относительное имя папки работы.
        """
        type_prefix = {
            "Практическая": "ПР",
            "Лабораторная": "ЛР",
            "Курсовая": "КР",
            "Реферат": "РЕФ",
        }.get(task_type, "ПР")

        folder_name = sanitize_folder_name(f"{type_prefix}_{number:02d}_{title}")
        subject_dir = self.get_subject_dir(subject_folder)
        task_dir = subject_dir / folder_name

        task_dir.mkdir(parents=True, exist_ok=True)
        (task_dir / self.SUBFOLDER_TASK).mkdir(parents=True, exist_ok=True)
        (task_dir / self.SUBFOLDER_WORKSPACE).mkdir(parents=True, exist_ok=True)
        (task_dir / self.SUBFOLDER_SUBMISSION).mkdir(parents=True, exist_ok=True)

        return folder_name

    def get_task_dir(self, subject_folder: str, task_folder: str) -> Path:
        return self.root_path / subject_folder / task_folder

    def get_task_subfolders(self, subject_folder: str, task_folder: str) -> Dict[str, Path]:
        """Возвращает пути к 3 основным подпапкам задания."""
        task_dir = self.get_task_dir(subject_folder, task_folder)
        return {
            "task": task_dir / self.SUBFOLDER_TASK,
            "workspace": task_dir / self.SUBFOLDER_WORKSPACE,
            "submission": task_dir / self.SUBFOLDER_SUBMISSION,
        }

    def list_task_files(
        self, subject_folder: str, task_folder: str
    ) -> Dict[str, List[Path]]:
        """Возвращает списки файлов по секциям: task, workspace, submission."""
        subfolders = self.get_task_subfolders(subject_folder, task_folder)
        result: Dict[str, List[Path]] = {}
        for key, folder in subfolders.items():
            if folder.exists() and folder.is_dir():
                try:
                    result[key] = sorted(
                        [p for p in folder.iterdir() if p.is_file()],
                        key=lambda p: p.name.lower(),
                    )
                except Exception:
                    result[key] = []
            else:
                result[key] = []
        return result

    # --- Запуск и интеграция с ОС Windows ---

    @staticmethod
    def open_path_in_explorer(path: Path) -> bool:
        """Открывает папку или выделяет файл в Проводнике Windows."""
        try:
            target = Path(path).resolve()
            if not target.exists():
                return False
            if target.is_dir():
                os.startfile(str(target))
            else:
                # Открыть проводник с выделенным файлом
                subprocess.Popen(["explorer", f"/select,{str(target)}"])
            return True
        except Exception:
            return False

    @staticmethod
    def format_file_size(size_bytes: int) -> str:
        """Форматирует размер файла в удобный для восприятия вид (Б, КБ, МБ)."""
        if size_bytes < 1024:
            return f"{size_bytes} Б"
        elif size_bytes < 1024 * 1024:
            return f"{size_bytes / 1024:.1f} КБ"
        else:
            return f"{size_bytes / (1024 * 1024):.1f} МБ"

    @staticmethod
    def get_file_icon_name(file_path: Path | str) -> str:
        """Возвращает название иконки Material Design по расширению файла."""
        suffix = Path(file_path).suffix.lower()
        if suffix in (".pdf",):
            return "picture_as_pdf"
        elif suffix in (".py", ".cpp", ".c", ".cs", ".java", ".js", ".ts", ".html", ".css", ".sql", ".sh"):
            return "code"
        elif suffix in (".zip", ".rar", ".7z", ".tar", ".gz"):
            return "archive"
        elif suffix in (".docx", ".doc", ".odt", ".rtf", ".txt", ".md"):
            return "description"
        else:
            return "insert_drive_file"

    @staticmethod
    def copy_file_to_subfolder(source_path: Path | str, target_folder: Path | str) -> Optional[Path]:
        """Копирует внешний файл в целевую директорию секции."""
        try:
            src = Path(source_path).resolve()
            dst_dir = Path(target_folder).resolve()
            if not src.is_file():
                return None
            dst_dir.mkdir(parents=True, exist_ok=True)
            dst_file = dst_dir / src.name
            shutil.copy2(src, dst_file)
            return dst_file
        except Exception:
            return None

    @staticmethod
    def delete_file(file_path: Path | str) -> bool:
        """Удаляет файл из директории."""
        try:
            p = Path(file_path).resolve()
            if p.is_file():
                p.unlink()
                return True
            return False
        except Exception:
            return False

    @staticmethod
    def open_file_default(file_path: Path) -> bool:
        """Открывает файл в стандартной программе Windows (os.startfile)."""
        try:
            target = Path(file_path).resolve()
            if target.is_file():
                os.startfile(str(target))
                return True
            return False
        except Exception:
            return False

    @staticmethod
    def open_in_vscode(path: Path) -> bool:
        """Открывает папку или файл в VS Code, если он установлен в PATH."""
        try:
            target = Path(path).resolve()
            subprocess.Popen(["code", str(target)], shell=True)
            return True
        except Exception:
            return False

    # --- Архивация и шаблоны (Шаг 3 и Шаг 4) ---

    def ensure_default_template(self) -> Path:
        """Создает дефолтный шаблон отчета в _Templates, если там нет ни одного docx."""
        templates_dir = self.root_path / self.FOLDER_TEMPLATES
        templates_dir.mkdir(parents=True, exist_ok=True)
        docx_files = list(templates_dir.glob("*.docx"))
        if docx_files:
            return docx_files[0]

        default_tpl = templates_dir / "Шаблон_отчета.docx"
        try:
            import docx
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Inches, Pt

            doc = docx.Document()
            for section in doc.sections:
                section.top_margin = Inches(0.79)
                section.bottom_margin = Inches(0.79)
                section.left_margin = Inches(1.18)
                section.right_margin = Inches(0.59)

            p_top = doc.add_paragraph("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РФ")
            p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_top.runs[0].font.size = Pt(12)
            p_top.runs[0].font.bold = True

            p_dept = doc.add_paragraph("Кафедра информационных и вычислительных систем")
            p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_dept.runs[0].font.size = Pt(12)

            for _ in range(4):
                doc.add_paragraph()

            p_title = doc.add_paragraph("ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ")
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_title.runs[0].font.size = Pt(18)
            p_title.runs[0].font.bold = True

            p_subj = doc.add_paragraph("по дисциплине: «{{SUBJECT}}»")
            p_subj.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_subj.runs[0].font.size = Pt(14)

            p_theme = doc.add_paragraph("Тема: «{{TASK_TYPE}} №{{TASK_NUM}}. {{TASK_TITLE}}»")
            p_theme.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_theme.runs[0].font.size = Pt(13)

            p_var = doc.add_paragraph("Вариант: {{VARIANT}}")
            p_var.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_var.runs[0].font.size = Pt(12)

            for _ in range(4):
                doc.add_paragraph()

            p_info = doc.add_paragraph()
            p_info.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r1 = p_info.add_run("Выполнил: студент {{STUDENT_NAME}}\nГруппа: {{GROUP}}\n")
            r1.font.size = Pt(12)

            for _ in range(4):
                doc.add_paragraph()

            p_bot = doc.add_paragraph("2026 г.")
            p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_bot.runs[0].font.size = Pt(12)

            doc.add_page_break()
            doc.add_heading("1. Цель работы", level=1)
            doc.add_paragraph("Изучить теоретические основы и получить практические навыки решения задач...")

            doc.add_heading("2. Ход выполнения работы", level=1)
            doc.add_paragraph("В процессе работы были выполнены следующие шаги...")

            doc.add_heading("3. Заключение и выводы", level=1)
            doc.add_paragraph("На основе полученных результатов сделаны выводы...")

            doc.save(str(default_tpl))
        except Exception:
            pass

        return default_tpl

    def generate_task_report(
        self,
        subject_name: str,
        task_number: int,
        task_type: str,
        task_title: str,
        student_name: str,
        student_group: str,
        subject_folder: str,
        task_folder: str,
        variant: Optional[str] = None,
    ) -> Path:
        """
        Создает отчет по шаблону из _Templates в 03_Submission,
        заменяя плейсхолдеры на реальные данные, и переименовывает файл
        в Отчет_[ПР3]_[Фамилия]_[Группа].docx.
        """
        import docx

        template_path = self.ensure_default_template()
        doc = docx.Document(str(template_path))

        type_prefix = {
            "Практическая": "ПР",
            "Лабораторная": "ЛР",
            "Курсовая": "КР",
            "Реферат": "РЕФ",
        }.get(task_type, "ПР")

        replacements = {
            "{{SUBJECT}}": subject_name or "Дисциплина",
            "{{TASK_NUM}}": str(task_number),
            "{{TASK_TYPE}}": task_type or "Практическая работа",
            "{{TASK_TITLE}}": task_title or "Тема работы",
            "{{STUDENT_NAME}}": student_name or "Студент",
            "{{GROUP}}": student_group or "Группа",
            "{{VARIANT}}": variant or "Без варианта",
        }

        def replace_in_paragraphs(paragraphs):
            for p in paragraphs:
                full_text = p.text
                if not any(k in full_text for k in replacements):
                    continue
                for run in p.runs:
                    for k, v in replacements.items():
                        if k in run.text:
                            run.text = run.text.replace(k, v)
                for k, v in replacements.items():
                    if k in p.text:
                        p.text = p.text.replace(k, v)

        replace_in_paragraphs(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    replace_in_paragraphs(cell.paragraphs)
        for section in doc.sections:
            replace_in_paragraphs(section.header.paragraphs)
            replace_in_paragraphs(section.footer.paragraphs)

        surname = student_name.split()[0] if student_name.strip() else "Студент"
        grp = student_group.replace(" ", "_") if student_group.strip() else "Группа"
        base_name = f"Отчет_{type_prefix}{task_number}_{surname}_{grp}"
        clean_name = f"{sanitize_folder_name(base_name, 60)}.docx"

        subfolders = self.get_task_subfolders(subject_folder, task_folder)
        sub_dir = subfolders["submission"]
        sub_dir.mkdir(parents=True, exist_ok=True)
        out_path = sub_dir / clean_name
        doc.save(str(out_path))
        return out_path

    def create_submission_zip(
        self, subject_folder: str, task_folder: str, zip_filename: Optional[str] = None
    ) -> Optional[Path]:
        """Упаковывает содержимое 03_Submission в zip-архив."""
        task_dir = self.get_task_dir(subject_folder, task_folder)
        submission_dir = task_dir / self.SUBFOLDER_SUBMISSION
        if not submission_dir.exists():
            return None

        if not zip_filename:
            zip_filename = f"{task_folder}_submission.zip"

        zip_path = submission_dir / zip_filename
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
            for item in submission_dir.rglob("*"):
                if item.is_file() and item.name != zip_filename and not item.name.endswith(".zip"):
                    arcname = item.relative_to(submission_dir)
                    zipf.write(item, arcname)
        return zip_path
